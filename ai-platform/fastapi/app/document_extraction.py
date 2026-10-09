import io
import zipfile
from http import HTTPStatus
from typing import Literal

from fastapi import HTTPException
from pydantic import BaseModel

from app.contract_intelligence import MAX_CONTRACT_CHARACTERS

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
PDF_MAGIC = b"%PDF-"
ZIP_MAGIC = b"PK\x03\x04"


class ExtractedText(BaseModel):
    format: Literal["pdf", "docx"]
    text: str
    truncated: bool
    page_count: int | None = None


def _reject(detail: str, code: int = HTTPStatus.UNPROCESSABLE_ENTITY) -> HTTPException:
    return HTTPException(status_code=code, detail=detail)


def extract_text(content: bytes) -> ExtractedText:
    """Detect the format from file signatures, never from the client-supplied name or type."""
    if len(content) > MAX_UPLOAD_BYTES:
        raise _reject("Document exceeds the size limit", HTTPStatus.REQUEST_ENTITY_TOO_LARGE)
    if content.startswith(PDF_MAGIC):
        return _from_pdf(content)
    if content.startswith(ZIP_MAGIC):
        return _from_docx(content)
    raise _reject("Only PDF and DOCX documents are supported")


def _bounded(text: str) -> tuple[str, bool]:
    text = text.replace("\x00", "").strip()
    if not text:
        raise _reject("No extractable text was found; scanned documents are not supported")
    if len(text) > MAX_CONTRACT_CHARACTERS:
        return text[:MAX_CONTRACT_CHARACTERS], True
    return text, False


def _from_pdf(content: bytes) -> ExtractedText:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise _reject("Encrypted documents are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise _reject("Document has too many pages")
        parts: list[str] = []
        total = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            parts.append(page_text)
            total += len(page_text)
            if total > MAX_CONTRACT_CHARACTERS:
                break
        page_count = len(reader.pages)
    except HTTPException:
        raise
    except Exception as exception:
        raise _reject("The PDF could not be read") from exception
    text, truncated = _bounded("\n".join(parts))
    return ExtractedText(format="pdf", text=text, truncated=truncated, page_count=page_count)


def _from_docx(content: bytes) -> ExtractedText:
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names:
                raise _reject("Only PDF and DOCX documents are supported")
            if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise _reject("Document expands beyond the allowed size")
        document = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
    except HTTPException:
        raise
    except Exception as exception:
        raise _reject("The DOCX could not be read") from exception
    text, truncated = _bounded("\n".join(parts))
    return ExtractedText(format="docx", text=text, truncated=truncated)
